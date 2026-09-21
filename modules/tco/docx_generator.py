import io
import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def gerar_docx_oficio(num_oficio, destinatario_nome, destinatario_cargo, orgao_destino, lista_materiais, pa_oficio, corpo_texto, emissor_nome, emissor_cargo, emissor_unidade):
    """Gera a versão editável em DOCX (Microsoft Word) do Ofício de Encaminhamento."""
    doc = Document()
    
    # Configuração de Margens (2 cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Cabeçalho Institucional
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_hdr1 = p_hdr.add_run("POLÍCIA MILITAR DE MINAS GERAIS\n")
    run_hdr1.bold = True
    run_hdr1.font.size = Pt(11)
    
    run_hdr2 = p_hdr.add_run(f"{emissor_unidade.upper()}\n")
    run_hdr2.bold = True
    run_hdr2.font.size = Pt(10)
    
    run_hdr3 = p_hdr.add_run("SEÇÃO DE CUSTÓDIA DE MATERIAIS E TCO - CREDS\n")
    run_hdr3.bold = True
    run_hdr3.font.size = Pt(9)
    run_hdr3.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph("─" * 55).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadados do Ofício
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"OFÍCIO Nº: {num_oficio}\n").bold = True
    p_meta.add_run(f"REF. P.A. / PROTOCOLO: {pa_oficio if pa_oficio else 'N/A'}\n").bold = True
    p_meta.add_run(f"DATA DE EMISSÃO: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}\n").bold = True

    # Destinatário
    p_dest = doc.add_paragraph()
    p_dest.add_run("Ao(À) Excelentíssimo(a) Senhor(a):\n").bold = True
    p_dest.add_run(f"{destinatario_nome.upper()}\n").bold = True
    p_dest.add_run(f"{destinatario_cargo.upper()}\n")
    p_dest.add_run(f"{orgao_destino.upper()}\n").bold = True

    # Título da Tabela
    p_tit_tb = doc.add_paragraph()
    run_tit = p_tit_tb.add_run("RELAÇÃO DE MATERIAIS ENCAMINHADOS (CADEIA DE CUSTÓDIA)")
    run_tit.bold = True
    run_tit.font.size = Pt(10)

    # Tabela de Materiais
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Nº REDS'
    hdr_cells[1].text = 'DESCRIÇÃO DO MATERIAL / LACRE'
    hdr_cells[2].text = 'NOME DO AUTOR'
    
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(9)

    for item in lista_materiais:
        row_cells = table.add_row().cells
        row_cells[0].text = str(item.get("num_reds", "N/I"))
        row_cells[1].text = f"{item.get('descricao', 'N/I')} (Qtd: {item.get('quantidade', '1.0')} {item.get('unidade_medida', 'UN')}) | Lacre: {item.get('involucro_lacre', 'N/I')}"
        row_cells[2].text = str(item.get("autores", "AUTOR NÃO INFORMADO"))

    doc.add_paragraph()

    # Corpo do Texto / Expediente
    p_corpo_tit = doc.add_paragraph()
    p_corpo_tit.add_run("TEOR DA SOLICITAÇÃO / HISTÓRICO:").bold = True
    
    p_corpo = doc.add_paragraph(corpo_texto)
    p_corpo.paragraph_format.line_spacing = 1.15
    p_corpo.paragraph_format.space_after = Pt(12)

    # Assinatura
    doc.add_paragraph("\n")
    p_ass = doc.add_paragraph()
    p_ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass.add_run("____________________________________________________\n")
    p_ass.add_run(f"{emissor_nome.upper()}\n").bold = True
    p_ass.add_run(f"{emissor_cargo} - {emissor_unidade}\n")

    # Guardar em memória BytesIO
    buffer_docx = io.BytesIO()
    doc.save(buffer_docx)
    buffer_docx.seek(0)
    return buffer_docx.getvalue()